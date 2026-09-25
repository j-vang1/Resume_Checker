"""Unit tests for resume matching and language detection."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document

from resume_matcher.language import detect_language
from resume_matcher.matching import extract_keywords, score_resume, score_resumes
from resume_matcher.parsers import extract_text_from_bytes, extract_text_from_docx


JOB_EN = """
We are hiring a Senior Python Software Engineer with experience in FastAPI,
Django, PostgreSQL, Docker, Kubernetes, and AWS. Candidates should know
REST APIs, microservices, CI/CD, pytest, and distributed systems.
"""

RESUME_STRONG = """
Jane Doe — Senior Software Engineer
Experienced Python developer specializing in FastAPI and Django.
Built microservices on AWS with Docker and Kubernetes.
Strong with PostgreSQL, REST APIs, CI/CD pipelines, and pytest.
"""

RESUME_WEAK = """
John Smith — Marketing Coordinator
Managed social media campaigns, wrote newsletters, organized events,
and coordinated brand partnerships for retail clients.
"""

JOB_ES = """
Buscamos un ingeniero de software con experiencia en Python, Django,
bases de datos PostgreSQL y despliegue en la nube con Docker.
"""


def test_detect_english_job_description():
    result = detect_language(JOB_EN)
    assert result["code"] == "en"
    assert result["name"] == "English"
    assert result["confidence"] > 0.5


def test_detect_spanish_job_description():
    result = detect_language(JOB_ES)
    assert result["code"] == "es"
    assert result["confidence"] > 0.5


def test_detect_too_short():
    result = detect_language("hi")
    assert result["code"] == "unknown"


def test_strips_lead_in_verbs_for_matching():
    """JD 'Review/Demonstrate X' should align with resume 'Led/Demonstrated X'."""
    from resume_matcher.matching import extract_keywords, score_resume, strip_lead_ins

    assert "thermal plunger" in strip_lead_ins(
        "Demonstrated thermal plunger design for ATE sockets"
    ).lower()
    assert "root cause" in strip_lead_ins(
        "Responsible for reviewing root cause analysis on failures"
    ).lower()
    assert strip_lead_ins("Review and approve test socket designs").lower().startswith(
        "test socket"
    )

    job = """
    Review and approve the design of test socket and thermal plunger hardware.
    Demonstrate root cause analysis on intermittent ATE failures.
    Provide DOE for thermal stability characterization.
    """
    resume = """
    - Led design of test socket and thermal plunger hardware for production ATE.
    - Demonstrated root cause analysis on intermittent ATE failures across sites.
    - Performed DOE for thermal stability characterization at -40C.
    """
    result = score_resume(job, resume, filename="ate.pdf", threshold=50)
    assert "socket" in result.matched_keywords or "thermal" in result.matched_keywords
    assert "plunger" in result.matched_keywords or "doe" in result.matched_keywords
    assert result.match_percent >= 40

    kws = extract_keywords(job)
    assert "review" not in kws
    assert "demonstrate" not in kws
    assert "provide" not in kws


def test_no_false_substring_keyword_matches():
    """'ate' must not match inside 'evaluate'; 'test' must not match 'latest'."""
    from resume_matcher.matching import _keyword_overlap, score_resume

    resume = """
    Experience
    - Evaluated candidates and created the latest marketing campaigns.
    - Related state gate analysis for process improvement.
    """
    # Force keywords that commonly false-hit via substring
    overlap, matched, missing = _keyword_overlap(
        ["ate", "test", "socket", "thermal", "plunger"],
        resume,
    )
    assert "ate" not in matched
    assert "test" not in matched
    assert "socket" not in matched
    assert "thermal" not in matched

    # Real whole-word hits still count
    resume2 = """
    - Debugged ATE handler issues and designed test socket hardware.
    - Thermal plunger characterization using DOE.
    """
    _, matched2, _ = _keyword_overlap(
        ["ate", "test", "socket", "thermal", "plunger", "doe"],
        resume2,
    )
    assert "ate" in matched2
    assert "socket" in matched2
    assert "thermal" in matched2
    assert "plunger" in matched2
    assert "doe" in matched2

    result = score_resume(
        "Need ATE test socket and thermal plunger experience.",
        resume,
        filename="false.pdf",
        threshold=50,
    )
    assert "ate" not in result.matched_keywords
    assert "socket" not in result.matched_keywords
    result = score_resume(JOB_EN, RESUME_STRONG, filename="strong.pdf", threshold=50)
    assert result.match_percent >= 50
    assert result.greenlit is True
    assert "python" in result.matched_keywords


def test_weak_resume_rejected_at_50():
    result = score_resume(JOB_EN, RESUME_WEAK, filename="weak.pdf", threshold=50)
    assert result.match_percent < 50
    assert result.greenlit is False


def test_threshold_is_tunable():
    result = score_resume(JOB_EN, RESUME_STRONG, filename="strong.pdf", threshold=99)
    # Even a strong resume may not hit 99%; assert threshold logic works
    assert result.greenlit == (result.match_percent >= 99)


def test_score_resumes_sorted_and_filtered():
    results = score_resumes(
        JOB_EN,
        [("strong.pdf", RESUME_STRONG), ("weak.pdf", RESUME_WEAK)],
        threshold=50,
    )
    assert len(results) == 2
    assert results[0].filename == "strong.pdf"
    assert results[0].greenlit is True
    assert results[1].filename == "weak.pdf"
    assert results[1].greenlit is False


def test_empty_inputs():
    empty_job = score_resume("", RESUME_STRONG, filename="a.pdf")
    assert empty_job.error
    assert empty_job.greenlit is False

    empty_resume = score_resume(JOB_EN, "", filename="b.pdf")
    assert empty_resume.error
    assert empty_resume.greenlit is False


def test_extract_keywords_finds_skills():
    keywords = extract_keywords(JOB_EN)
    joined = " ".join(keywords)
    assert "python" in joined
    assert "fastapi" in joined or "django" in joined


def test_docx_parser_roundtrip():
    document = Document()
    document.add_paragraph("Alice Example")
    document.add_paragraph("Skills: Python, FastAPI, PostgreSQL")
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text_from_docx(buffer.getvalue())
    assert "Alice Example" in text
    assert "Python" in text


def test_extract_text_from_bytes_docx():
    document = Document()
    document.add_paragraph(
        "Resume text for parsing with enough characters to pass the "
        "minimum extractable-text check used for designed PDFs."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text_from_bytes(buffer.getvalue(), "resume.docx")
    assert "Resume text" in text


def test_unsupported_extension():
    with pytest.raises(Exception):
        extract_text_from_bytes(b"not a resume", "resume.txt")
