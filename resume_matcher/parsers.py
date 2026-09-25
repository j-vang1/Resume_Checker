"""Extract plain text from PDF and Word resume files."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc"}


class UnsupportedFileTypeError(ValueError):
    """Raised when a file type cannot be parsed."""


def _normalize_ext(filename: str) -> str:
    return Path(filename).suffix.lower()


def extract_text_from_pdf(data: bytes) -> str:
    """Extract text from a PDF byte stream."""
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n".join(pages).strip()


def extract_text_from_docx(data: bytes) -> str:
    """Extract text from a DOCX byte stream."""
    document = Document(io.BytesIO(data))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    """
    Extract text from uploaded file bytes based on filename extension.

    Supports ``.pdf`` and ``.docx``. Legacy ``.doc`` is not supported;
    convert those to DOCX or PDF first.
    """
    ext = _normalize_ext(filename)
    if ext == ".pdf":
        return extract_text_from_pdf(data)
    if ext == ".docx":
        return extract_text_from_docx(data)
    if ext == ".doc":
        raise UnsupportedFileTypeError(
            f"'{filename}' is a legacy .doc file. Please convert it to .docx or .pdf."
        )
    raise UnsupportedFileTypeError(
        f"Unsupported file type '{ext}' for '{filename}'. "
        f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS - {'.doc'}))}."
    )


def extract_text_from_path(path: str | Path) -> str:
    """Extract text from a file on disk."""
    path = Path(path)
    return extract_text_from_bytes(path.read_bytes(), path.name)